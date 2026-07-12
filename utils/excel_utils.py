from openpyxl import load_workbook
from datetime import datetime


class ExcelUtils:

    def __init__(self, file_path):
        self.file_path = file_path
        self.workbook = load_workbook(file_path)
        self.sheet = self.workbook["LoginData"]

    # Total rows
    def get_row_count(self):
        return self.sheet.max_row

    # Read username
    def get_username(self, row):
        return self.sheet.cell(row=row, column=2).value

    # Read password
    def get_password(self, row):
        return self.sheet.cell(row=row, column=3).value

    # Write Pass/Fail
    def write_result(self, row, result):
        self.sheet.cell(row=row, column=7).value = result

    # Write Date
    def write_date(self, row):
        self.sheet.cell(row=row, column=4).value = datetime.now().strftime("%d-%m-%Y")

    # Write Time
    def write_time(self, row):
        self.sheet.cell(row=row, column=5).value = datetime.now().strftime("%H:%M:%S")

    # Save Excel
    def save(self):
        self.workbook.save(self.file_path)
